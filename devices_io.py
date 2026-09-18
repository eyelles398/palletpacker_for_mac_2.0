"""
Импорт/экспорт справочника устройств (devices.json) в Excel.

Формат файла:
    A — Название
    B — Тип ("обычная" / "МАСТЕР")
    C — Секция
    D — Ш, мм
    E — Г, мм
    F — В, мм
    G — Вес, кг (может быть пустым)
"""
from __future__ import annotations

import openpyxl


EXPORT_HEADERS = ["Название", "Тип", "Секция",
                  "Ш, мм", "Г, мм", "В, мм", "Вес, кг"]


# ---------- утилиты ----------

def _to_float(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", ".").strip())
    except (ValueError, TypeError):
        return None


def _safe_int(v, default=0):
    try:
        return int(round(float(v)))
    except (TypeError, ValueError):
        return default


# ---------- экспорт ----------

def export_devices_xlsx(devices, path):
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Справочник"

    title_font = Font(bold=True, size=14, color="2E7D32")
    head_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="E8F5E9")
    small_font = Font(size=10, color="666666")

    ws["A1"] = "Справочник устройств"
    ws["A1"].font = title_font
    ws.merge_cells("A1:G1")
    ws["A2"] = f"Записей: {len(devices)}"
    ws["A2"].font = small_font

    for c, h in enumerate(EXPORT_HEADERS, 1):
        cell = ws.cell(row=4, column=c, value=h)
        cell.font = head_font
        cell.fill = header_fill

    row = 5
    for d in devices:
        name = d.get("name", "")
        kind = "МАСТЕР" if d.get("type") == "master" else "обычная"
        section = d.get("section", "") or ""
        dims = d.get("dims_mm") or (0, 0, 0)
        w = d.get("weight_kg")
        weight = "" if w is None else float(w)

        ws.cell(row=row, column=1, value=name)
        ws.cell(row=row, column=2, value=kind)
        ws.cell(row=row, column=3, value=section)
        ws.cell(row=row, column=4, value=_safe_int(dims[0]))
        ws.cell(row=row, column=5, value=_safe_int(dims[1]))
        ws.cell(row=row, column=6, value=_safe_int(dims[2]))
        ws.cell(row=row, column=7, value=weight)
        row += 1

    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 32
    for c in "DEFG":
        ws.column_dimensions[c].width = 11

    wb.save(path)


# ---------- импорт ----------

def import_devices_xlsx(path):
    """
    Читает Excel и возвращает (devices, errors):
      devices — список dict'ов формата devices.json
      errors  — список строк с описаниями проблемных строк
    """
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()

    devices = []
    errors = []

    # ищем строку-заголовок (если есть)
    header_row = None
    for i, row in enumerate(rows):
        if not row:
            continue
        a = row[0] if len(row) > 0 else None
        if isinstance(a, str):
            a_low = a.strip().lower()
            if a_low in ("название", "модель", "name", "наименование"):
                header_row = i
                break

    start = (header_row + 1) if header_row is not None else 0

    for i in range(start, len(rows)):
        row = rows[i]
        if not row:
            continue

        a = row[0] if len(row) > 0 else None
        b = row[1] if len(row) > 1 else None
        c = row[2] if len(row) > 2 else None
        d = row[3] if len(row) > 3 else None
        e = row[4] if len(row) > 4 else None
        f = row[5] if len(row) > 5 else None
        g = row[6] if len(row) > 6 else None

        if a is None or not str(a).strip():
            continue

        name = str(a).strip()
        low = name.lower()
        # пропускаем служебные строки
        if low.startswith(("справочник", "итого", "total")):
            continue

        kind_raw = str(b).strip().lower() if b is not None else ""
        kind = "master" if ("мастер" in kind_raw or kind_raw == "master") else "box"
        section = str(c).strip() if c is not None else ""

        w = _to_float(d)
        dd = _to_float(e)
        h = _to_float(f)
        weight = _to_float(g)

        if w is None or dd is None or h is None:
            errors.append(f"Строка {i + 1}: нет размеров для «{name}»")
            continue

        devices.append({
            "name": name,
            "section": section,
            "type": kind,
            "dims_mm": [w, dd, h],
            "weight_kg": weight,
        })

    return devices, errors