"""
Экспорт результатов расчёта в Excel и PDF.
Поддерживает учёт веса тары (брутто), устойчивость (ЦТ),
индивидуальные параметры паллет и встраивание 3D-скриншота в PDF.
"""
import os
import tempfile

from PySide6.QtPrintSupport import QPrinter
from PySide6.QtGui import QTextDocument, QPageSize
from PySide6.QtCore import QMarginsF, QUrl


# ============================================================
#                       PDF
# ============================================================

def export_pdf(data, path, image=None):
    """
    data  — словарь с результатами расчёта (см. main.py)
    path  — куда сохранить PDF
    image — QImage со скриншотом 3D-вида (необязательно)
    """
    tmp_png = None
    img_src = None

    # QPainter внутри QPrinter надёжно подхватывает только
    # картинки из локальных файлов (а не из QTextDocument resources)
    if image is not None:
        try:
            fd, tmp_png = tempfile.mkstemp(suffix=".png",
                                           prefix="pp3d_")
            os.close(fd)
            if not image.save(tmp_png, "PNG"):
                tmp_png = None
        except Exception:
            tmp_png = None

    if tmp_png:
        # подставляем как file:/// URL — так QTextDocument корректно
        # отрисует картинку при печати через QPrinter
        img_src = QUrl.fromLocalFile(tmp_png).toString()

    html = _build_html(data, img_src=img_src)

    doc = QTextDocument()
    doc.setHtml(html)

    printer = QPrinter(QPrinter.HighResolution)
    printer.setOutputFormat(QPrinter.PdfFormat)
    printer.setOutputFileName(path)
    printer.setPageSize(QPageSize(QPageSize.A4))
    printer.setPageMargins(QMarginsF(15, 15, 15, 15))

    doc.print_(printer)

    # удаляем временный файл
    if tmp_png and os.path.exists(tmp_png):
        try:
            os.remove(tmp_png)
        except Exception:
            pass


def _build_html(data, img_src=None):
    p = data["pallet"]
    pallets = data["pallets"]
    unpacked = data["unpacked"]
    by_model = data["by_model"]

    use_tare = bool(p.get("use_tare", False))
    tare = float(p.get("tare_weight", 0.0))
    uniform = bool(p.get("uniform", True))

    goods_weight = sum(x["weight_kg"] for x in pallets)
    tare_total = tare * len(pallets) if use_tare else 0.0
    total_weight = goods_weight + tare_total

    total_volume = sum(x["volume_m3"] for x in pallets)
    total_bbox = sum(x["bbox_volume_m3"] for x in pallets)
    ordered_total = sum(it["qty"] for it in data["items"])
    placed_total = sum(len(x["placed"]) for x in pallets)

    css = """
    <style>
      body { font-family: Arial, sans-serif; font-size: 11pt; }
      h1 { color: #2e7d32; font-size: 18pt; margin-bottom: 4px; }
      h2 { color: #333; font-size: 14pt; margin-top: 14px; }
      table { border-collapse: collapse; margin: 6px 0; width: 100%; }
      th, td { border: 1px solid #888; padding: 4px 6px; font-size: 9pt; }
      th { background: #e8f5e9; text-align: left; }
      .num { text-align: right; }
      .warn { color: #b71c1c; font-weight: bold; }
      .small { color: #666; font-size: 9pt; }
      .total { font-weight: bold; background: #f1f8e9; }
      .img3d { margin-top: 8px; }
    </style>
    """
    h = [css]
    h.append("<h1>Отчёт об укладке коробок на паллету</h1>")

    if uniform:
        h.append(
            f'<p class="small">Паллета: '
            f'{p["L"]:.0f}×{p["W"]:.0f}×{p["H"]:.0f} мм, '
            f'макс. вес {p["max_weight"]:.0f} кг</p>'
        )
    else:
        h.append(
            f'<p class="small">Паллеты <b>индивидуальные</b> '
            f'(базовые: {p["L"]:.0f}×{p["W"]:.0f}×{p["H"]:.0f} мм, '
            f'до {p["max_weight"]:.0f} кг)</p>'
        )

    # Сводка
    h.append("<h2>Сводка</h2>")
    h.append("<table>")
    h.append(f'<tr><td>Всего физических паллет:</td><td class="num">{len(pallets)}</td></tr>')
    h.append(f'<tr><td>Уложено коробок:</td><td class="num">{placed_total} из {ordered_total}</td></tr>')
    h.append(f'<tr><td>Вес коробок:</td><td class="num">{goods_weight:.2f} кг</td></tr>')
    if use_tare:
        h.append(
            f'<tr><td>Вес тары ({tare:.1f} кг × {len(pallets)} пал.):</td>'
            f'<td class="num">{tare_total:.2f} кг</td></tr>'
        )
        h.append(
            f'<tr class="total"><td>ОБЩИЙ ВЕС БРУТТО:</td>'
            f'<td class="num">{total_weight:.2f} кг</td></tr>'
        )
    else:
        h.append(
            f'<tr class="total"><td>ОБЩИЙ ВЕС:</td>'
            f'<td class="num">{total_weight:.2f} кг</td></tr>'
        )
    h.append(f'<tr><td>Объём коробок:</td><td class="num">{total_volume:.4f} м³</td></tr>')
    h.append(f'<tr><td>Габаритный объём:</td><td class="num">{total_bbox:.4f} м³</td></tr>')
    h.append("</table>")

    if unpacked:
        h.append(f'<p class="warn">⚠ Не поместилось: {len(unpacked)} шт.</p>')

    unstable = [x for x in pallets if not x.get("cog_stable", True)]
    if unstable:
        h.append(
            f'<p class="warn">⚠ ЦТ вне безопасной зоны '
            f'(±25% от центра) у {len(unstable)} паллет(ы).</p>'
        )

    # 3D-визуализация
    if img_src:
        h.append("<h2>Визуализация укладки (3D)</h2>")
        h.append('<p class="small">Скриншот 3D-вида с текущего расчёта.</p>')
        h.append(f'<p><img class="img3d" src="{img_src}" width="680"></p>')

    # Паллеты
    h.append("<h2>Паллеты</h2>")
    h.append("<table>")
    if use_tare:
        h.append(
            "<tr>"
            "<th>№</th><th>Группа</th><th>Коробок</th>"
            "<th>Паллета, мм</th>"
            "<th>Вес коробок, кг</th><th>Тара, кг</th><th>Брутто, кг</th>"
            "<th>Объём коробок, м³</th>"
            "<th>Габаритный объём, м³</th><th>Высота, мм</th>"
            "<th>ЦТ X, мм</th><th>ЦТ Y, мм</th><th>Устойч.</th>"
            "</tr>"
        )
    else:
        h.append(
            "<tr>"
            "<th>№</th><th>Группа</th><th>Коробок</th>"
            "<th>Паллета, мм</th>"
            "<th>Вес, кг</th>"
            "<th>Объём коробок, м³</th>"
            "<th>Габаритный объём, м³</th><th>Высота, мм</th>"
            "<th>ЦТ X, мм</th><th>ЦТ Y, мм</th><th>Устойч.</th>"
            "</tr>"
        )

    for i, pal in enumerate(pallets, 1):
        bx, by, bz = pal["bbox_size_mm"]
        ox, oy = pal.get("cog_offset_mm", (0.0, 0.0))
        stable = pal.get("cog_stable", True)
        stable_str = "✓" if stable else "<span class='warn'>⚠ НЕТ</span>"
        pL = pal.get("L", p["L"])
        pW = pal.get("W", p["W"])
        pH = pal.get("H", p["H"])
        pallet_dims = f"{pL:.0f}×{pW:.0f}×{pH:.0f}"
        if use_tare:
            gross = pal["weight_kg"] + tare
            h.append(
                f"<tr>"
                f"<td class='num'>{i}</td>"
                f"<td class='num'>{pal['group']}</td>"
                f"<td class='num'>{len(pal['placed'])}</td>"
                f"<td class='num'>{pallet_dims}</td>"
                f"<td class='num'>{pal['weight_kg']:.2f}</td>"
                f"<td class='num'>{tare:.2f}</td>"
                f"<td class='num'>{gross:.2f}</td>"
                f"<td class='num'>{pal['volume_m3']:.4f}</td>"
                f"<td class='num'>{pal['bbox_volume_m3']:.4f}"
                f"  <span class='small'>({bx:.0f}×{by:.0f}×{bz:.0f})</span></td>"
                f"<td class='num'>{pal['height_mm']:.0f}</td>"
                f"<td class='num'>{ox:+.0f}</td>"
                f"<td class='num'>{oy:+.0f}</td>"
                f"<td class='num'>{stable_str}</td>"
                f"</tr>"
            )
        else:
            h.append(
                f"<tr>"
                f"<td class='num'>{i}</td>"
                f"<td class='num'>{pal['group']}</td>"
                f"<td class='num'>{len(pal['placed'])}</td>"
                f"<td class='num'>{pallet_dims}</td>"
                f"<td class='num'>{pal['weight_kg']:.2f}</td>"
                f"<td class='num'>{pal['volume_m3']:.4f}</td>"
                f"<td class='num'>{pal['bbox_volume_m3']:.4f}"
                f"  <span class='small'>({bx:.0f}×{by:.0f}×{bz:.0f})</span></td>"
                f"<td class='num'>{pal['height_mm']:.0f}</td>"
                f"<td class='num'>{ox:+.0f}</td>"
                f"<td class='num'>{oy:+.0f}</td>"
                f"<td class='num'>{stable_str}</td>"
                f"</tr>"
            )
    h.append("</table>")

    # По моделям
    h.append("<h2>По моделям</h2>")
    h.append("<table>")
    h.append(
        "<tr><th>Модель</th><th>Заказано</th>"
        "<th>Уложено</th><th>Не поместилось</th>"
        "<th>Размер (Ш×Г×В), мм</th><th>Вес, кг/шт</th></tr>"
    )
    for name, d in by_model.items():
        L, W, H = d["dims"]
        un = d["unpacked"]
        un_str = f"<span class='warn'>{un}</span>" if un else "0"
        h.append(
            f"<tr>"
            f"<td>{name}</td>"
            f"<td class='num'>{d['qty_ordered']}</td>"
            f"<td class='num'>{d['placed']}</td>"
            f"<td class='num'>{un_str}</td>"
            f"<td class='num'>{int(L)}×{int(W)}×{int(H)}</td>"
            f"<td class='num'>{d['weight_kg'] or 0}</td>"
            f"</tr>"
        )
    h.append("</table>")

    h.append('<p class="small">Отчёт сгенерирован программой «Расчёт укладки коробок на паллету».</p>')
    return "".join(h)


# ============================================================
#                       Excel
# ============================================================

def export_excel(data, path):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    p = data["pallet"]
    pallets = data["pallets"]
    unpacked = data["unpacked"]
    by_model = data["by_model"]

    use_tare = bool(p.get("use_tare", False))
    tare = float(p.get("tare_weight", 0.0))
    uniform = bool(p.get("uniform", True))

    goods_weight = sum(x["weight_kg"] for x in pallets)
    tare_total = tare * len(pallets) if use_tare else 0.0
    total_weight = goods_weight + tare_total

    total_volume = sum(x["volume_m3"] for x in pallets)
    total_bbox = sum(x["bbox_volume_m3"] for x in pallets)
    ordered_total = sum(it["qty"] for it in data["items"])
    placed_total = sum(len(x["placed"]) for x in pallets)

    wb = openpyxl.Workbook()

    title_font = Font(bold=True, size=14, color="2E7D32")
    head_font = Font(bold=True)
    total_font = Font(bold=True, color="1B5E20")
    warn_font = Font(bold=True, color="B71C1C")
    header_fill = PatternFill("solid", fgColor="E8F5E9")
    total_fill = PatternFill("solid", fgColor="F1F8E9")
    warn_fill = PatternFill("solid", fgColor="FFEBEE")

    # -------- Лист 1: Сводка --------
    ws = wb.active
    ws.title = "Сводка"

    ws["A1"] = "Отчёт об укладке коробок на паллету"
    ws["A1"].font = title_font
    ws.merge_cells("A1:C1")

    ws["A3"] = "Параметры паллеты (по умолчанию)"
    ws["A3"].font = head_font
    ws["A4"] = "Длина, мм";         ws["B4"] = p["L"]
    ws["A5"] = "Ширина, мм";        ws["B5"] = p["W"]
    ws["A6"] = "Высота, мм";        ws["B6"] = p["H"]
    ws["A7"] = "Макс. вес, кг";     ws["B7"] = p["max_weight"]
    ws["A8"] = "Вес тары, кг";      ws["B8"] = tare
    ws["A9"] = "Учитывать тару";    ws["B9"] = "да" if use_tare else "нет"
    ws["A10"] = "Паллеты одинаковые"; ws["B10"] = "да" if uniform else "НЕТ"

    ws["A12"] = "Итоги"; ws["A12"].font = head_font
    row = 13
    ws.cell(row=row, column=1, value="Всего паллет");              ws.cell(row=row, column=2, value=len(pallets)); row += 1
    ws.cell(row=row, column=1, value="Заказано коробок");          ws.cell(row=row, column=2, value=ordered_total); row += 1
    ws.cell(row=row, column=1, value="Уложено коробок");           ws.cell(row=row, column=2, value=placed_total); row += 1
    ws.cell(row=row, column=1, value="Не поместилось, шт.");       ws.cell(row=row, column=2, value=len(unpacked)); row += 1
    ws.cell(row=row, column=1, value="Вес коробок, кг");           ws.cell(row=row, column=2, value=round(goods_weight, 2)); row += 1
    if use_tare:
        ws.cell(row=row, column=1, value="Вес тары, кг").font = head_font
        ws.cell(row=row, column=2, value=round(tare_total, 2)).font = head_font
        row += 1
        ws.cell(row=row, column=1, value="ОБЩИЙ ВЕС БРУТТО, кг").font = total_font
        ws.cell(row=row, column=2, value=round(total_weight, 2)).font = total_font
        for c in (1, 2):
            ws.cell(row=row, column=c).fill = total_fill
        row += 1
    else:
        ws.cell(row=row, column=1, value="ОБЩИЙ ВЕС, кг").font = total_font
        ws.cell(row=row, column=2, value=round(total_weight, 2)).font = total_font
        for c in (1, 2):
            ws.cell(row=row, column=c).fill = total_fill
        row += 1
    ws.cell(row=row, column=1, value="Объём коробок, м³");         ws.cell(row=row, column=2, value=round(total_volume, 4)); row += 1
    ws.cell(row=row, column=1, value="Габаритный объём, м³");      ws.cell(row=row, column=2, value=round(total_bbox, 4)); row += 1

    unstable = [x for x in pallets if not x.get("cog_stable", True)]
    row += 1
    ws.cell(row=row, column=1, value="Неустойчивых паллет (ЦТ)").font = head_font
    c2 = ws.cell(row=row, column=2, value=len(unstable))
    if unstable:
        c2.font = warn_font
        c2.fill = warn_fill

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 16

    # -------- Лист 2: Паллеты --------
    ws2 = wb.create_sheet("Паллеты")
    if use_tare:
        headers = ["№", "Группа", "Продолжение", "Коробок",
                   "Паллета Д, мм", "Паллета Ш, мм", "Паллета В, мм",
                   "Вес коробок, кг", "Тара, кг", "Брутто, кг",
                   "Объём коробок, м³", "Габаритный объём, м³",
                   "Ширина bbox, мм", "Глубина bbox, мм", "Высота bbox, мм",
                   "Смещение ЦТ X, мм", "Смещение ЦТ Y, мм", "Устойчиво"]
    else:
        headers = ["№", "Группа", "Продолжение", "Коробок",
                   "Паллета Д, мм", "Паллета Ш, мм", "Паллета В, мм",
                   "Вес, кг",
                   "Объём коробок, м³", "Габаритный объём, м³",
                   "Ширина bbox, мм", "Глубина bbox, мм", "Высота bbox, мм",
                   "Смещение ЦТ X, мм", "Смещение ЦТ Y, мм", "Устойчиво"]
    for c, htxt in enumerate(headers, 1):
        cell = ws2.cell(row=1, column=c, value=htxt)
        cell.font = head_font
        cell.fill = header_fill

    for i, pal in enumerate(pallets, 1):
        bx, by, bz = pal["bbox_size_mm"]
        pL = pal.get("L", p["L"])
        pW = pal.get("W", p["W"])
        pH = pal.get("H", p["H"])
        row = i + 1
        col = 1
        ws2.cell(row=row, column=col, value=i); col += 1
        ws2.cell(row=row, column=col, value=pal["group"]); col += 1
        ws2.cell(row=row, column=col, value=pal["attempt"] - 1); col += 1
        ws2.cell(row=row, column=col, value=len(pal["placed"])); col += 1
        ws2.cell(row=row, column=col, value=round(pL, 1)); col += 1
        ws2.cell(row=row, column=col, value=round(pW, 1)); col += 1
        ws2.cell(row=row, column=col, value=round(pH, 1)); col += 1
        ws2.cell(row=row, column=col, value=round(pal["weight_kg"], 2)); col += 1
        if use_tare:
            ws2.cell(row=row, column=col, value=round(tare, 2)); col += 1
            ws2.cell(row=row, column=col, value=round(pal["weight_kg"] + tare, 2)); col += 1
        ws2.cell(row=row, column=col, value=round(pal["volume_m3"], 4)); col += 1
        ws2.cell(row=row, column=col, value=round(pal["bbox_volume_m3"], 4)); col += 1
        ws2.cell(row=row, column=col, value=round(bx, 1)); col += 1
        ws2.cell(row=row, column=col, value=round(by, 1)); col += 1
        ws2.cell(row=row, column=col, value=round(bz, 1)); col += 1

        ox, oy = pal.get("cog_offset_mm", (0.0, 0.0))
        stable = pal.get("cog_stable", True)
        ws2.cell(row=row, column=col, value=round(ox, 1)); col += 1
        ws2.cell(row=row, column=col, value=round(oy, 1)); col += 1
        c_st = ws2.cell(row=row, column=col,
                        value="да" if stable else "НЕТ")
        if not stable:
            c_st.font = warn_font
            c_st.fill = warn_fill

    for c in range(1, len(headers) + 1):
        ws2.column_dimensions[get_column_letter(c)].width = 16

    # -------- Лист 3: По моделям --------
    ws3 = wb.create_sheet("По моделям")
    headers3 = ["Модель", "Заказано", "Уложено",
                "Не поместилось", "Ш, мм", "Г, мм", "В, мм", "Вес, кг/шт"]
    for c, htxt in enumerate(headers3, 1):
        cell = ws3.cell(row=1, column=c, value=htxt)
        cell.font = head_font
        cell.fill = header_fill

    for i, (name, d) in enumerate(by_model.items(), 1):
        L, W, H = d["dims"]
        row = i + 1
        ws3.cell(row=row, column=1, value=name)
        ws3.cell(row=row, column=2, value=d["qty_ordered"])
        ws3.cell(row=row, column=3, value=d["placed"])
        ws3.cell(row=row, column=4, value=d["unpacked"])
        ws3.cell(row=row, column=5, value=int(L))
        ws3.cell(row=row, column=6, value=int(W))
        ws3.cell(row=row, column=7, value=int(H))
        ws3.cell(row=row, column=8, value=d["weight_kg"] or 0)

    for c in range(1, len(headers3) + 1):
        ws3.column_dimensions[get_column_letter(c)].width = 18

    # -------- Лист 4: Размещение (все коробки) --------
    ws4 = wb.create_sheet("Размещение")
    headers4 = ["№ паллеты", "Группа", "Модель",
                "X, мм", "Y, мм", "Z, мм",
                "Ш, мм", "Г, мм", "В, мм", "Вес, кг"]
    for c, htxt in enumerate(headers4, 1):
        cell = ws4.cell(row=1, column=c, value=htxt)
        cell.font = head_font
        cell.fill = header_fill

    row = 2
    for i, pal in enumerate(pallets, 1):
        for b in pal["placed"]:
            model = b["name"].split("#")[0]
            ws4.cell(row=row, column=1, value=i)
            ws4.cell(row=row, column=2, value=pal["group"])
            ws4.cell(row=row, column=3, value=model)
            ws4.cell(row=row, column=4, value=round(b["x"], 1))
            ws4.cell(row=row, column=5, value=round(b["y"], 1))
            ws4.cell(row=row, column=6, value=round(b["z"], 1))
            ws4.cell(row=row, column=7, value=round(b["dx"], 1))
            ws4.cell(row=row, column=8, value=round(b["dy"], 1))
            ws4.cell(row=row, column=9, value=round(b["dz"], 1))
            ws4.cell(row=row, column=10, value=round(b["weight"], 3))
            row += 1

    for c in range(1, len(headers4) + 1):
        ws4.column_dimensions[get_column_letter(c)].width = 14

    wb.save(path)