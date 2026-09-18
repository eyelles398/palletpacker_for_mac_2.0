"""
Экспорт отправок в PDF и Excel.
Excel: 3 листа — Отправки, Состав (расширенный), Сводка.
PDF для ТК: краткая таблица + вес + объём.
PDF внутренний: расширенная таблица с размерами и весом.
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from PySide6.QtPrintSupport import QPrinter
from PySide6.QtGui import QTextDocument, QPageSize
from PySide6.QtCore import QMarginsF

import shipments as shp


# ============================================================
#                    PDF для транспортной компании
# ============================================================

def export_tk_pdf(record: dict, path: str) -> None:
    html = _build_tk_html(record)
    _print_html(html, path)


def _build_tk_html(r: dict) -> str:
    entity_type = "Юридическое лицо" if r.get("entity_type") == "legal" \
        else "Физическое лицо"

    break_str = "—"
    if r.get("has_break"):
        break_str = r.get("break_time", "") or "есть"

    call_mark = "☑" if r.get("call_before_30min") else "☐"
    pass_mark = "☑" if r.get("need_pass") else "☐"

    address = (r.get("address", "") or "").replace("\n", "<br>")

    # состав с весом
    items_html = ""
    total_weight = 0.0
    for it in r.get("items", []):
        model = _esc(it.get("model", ""))
        qty = int(it.get("qty", 0))
        try:
            w1 = float(it.get("weight_kg", 0) or 0)
        except (TypeError, ValueError):
            w1 = 0.0
        w_total = qty * w1
        total_weight += w_total
        weight_str = f"{w1:g}" if w1 else "—"
        total_str = f"{w_total:.2f}" if w_total else "—"
        items_html += (
            f"<tr>"
            f"<td>{model}</td>"
            f"<td class='num'>{qty}</td>"
            f"<td class='num'>{weight_str}</td>"
            f"<td class='num'>{total_str}</td>"
            f"</tr>"
        )
    if not items_html:
        items_html = "<tr><td colspan='4' style='text-align:center;color:#888'>" \
                     "нет позиций</td></tr>"

    total_weight_row = (
        f"<tr><td colspan='3' style='text-align:right;font-weight:bold'>"
        f"ИТОГО вес:</td>"
        f"<td class='num' style='font-weight:bold'>{total_weight:.2f} кг</td>"
        f"</tr>"
    ) if total_weight else ""

    # грузовые характеристики (без паллет — они водителю не важны)
    try:
        volume = float(r.get("cargo_volume_m3", 0) or 0)
    except (TypeError, ValueError):
        volume = 0
    try:
        bbox = float(r.get("cargo_bbox_volume_m3", 0) or 0)
    except (TypeError, ValueError):
        bbox = 0

    cargo_rows = []
    if total_weight:
        cargo_rows.append(("Общий вес груза, кг", f"{total_weight:.2f}"))
    if volume:
        cargo_rows.append(("Объём коробок, м³", f"{volume:.4f}"))
    if bbox:
        cargo_rows.append(("Габаритный объём, м³", f"{bbox:.4f}"))

    cargo_html = ""
    if cargo_rows:
        cargo_html = ['<table class="cargo-table">']
        for k, v in cargo_rows:
            cargo_html.append(
                f'<tr><td>{_esc(k)}</td><td class="num">{_esc(v)}</td></tr>'
            )
        cargo_html.append("</table>")
        cargo_html = "".join(cargo_html)

    css = """
    <style>
      body { font-family: Arial, sans-serif; color: #111; }
      h1 { color: #1976d2; font-size: 22pt; margin: 0 0 4px 0; }
      h2 { font-size: 13pt; color: #333; margin: 14px 0 6px 0;
           border-bottom: 1px solid #ccc; padding-bottom: 2px; }
      .meta { color: #666; font-size: 10pt; margin-bottom: 14px; }
      .addr-box { border: 2px solid #1976d2; border-radius: 6px;
                  padding: 12px; background: #f5faff; }
      .addr-label { color: #1976d2; font-size: 11pt; font-weight: bold;
                    margin-bottom: 4px; }
      .addr-text { font-size: 16pt; font-weight: bold; line-height: 1.35; }
      table { border-collapse: collapse; width: 100%; margin-top: 6px; }
      td, th { border: 1px solid #888; padding: 5px 8px; font-size: 10pt;
               vertical-align: top; }
      th { background: #e3f2fd; text-align: left; }
      .num { text-align: right; }
      .kv { margin: 3px 0; font-size: 11pt; }
      .kv b { display: inline-block; min-width: 180px; }
      .check { font-size: 13pt; }
      .check .mark { font-size: 15pt; font-weight: bold;
                     margin-right: 6px; color: #2e7d32; }
      .comment { background: #fffde7; border-left: 4px solid #fbc02d;
                 padding: 8px 12px; margin-top: 6px; font-size: 10pt; }
      .cargo-table { width: auto; min-width: 55%; }
      .cargo-table td { padding: 4px 10px; font-size: 11pt; }
      .cargo-table td:first-child { background: #f5f5f5; font-weight: bold; }
    </style>
    """

    h = [css]
    h.append("<h1>ЗАДАНИЕ НА ДОСТАВКУ</h1>")
    h.append(
        f'<p class="meta">Отправка № <b>{_esc(r.get("id", ""))}</b>'
        f' &nbsp;·&nbsp; от {_esc(r.get("ship_date", ""))}</p>'
    )

    h.append('<div class="addr-box">')
    h.append('<div class="addr-label">АДРЕС ДОСТАВКИ</div>')
    h.append(f'<div class="addr-text">{address or "—"}</div>')
    h.append('</div>')

    h.append("<h2>Получатель</h2>")
    h.append(f'<div class="kv"><b>Тип лица:</b> {entity_type}</div>')
    h.append(f'<div class="kv"><b>Наименование / ФИО:</b> '
             f'{_esc(r.get("entity_name", "")) or "—"}</div>')
    if r.get("entity_type") == "legal":
        h.append(f'<div class="kv"><b>ИНН:</b> '
                 f'{_esc(r.get("inn", "")) or "—"}</div>')

    h.append("<h2>Контактные данные</h2>")
    h.append(f'<div class="kv"><b>Контактное лицо:</b> '
             f'{_esc(r.get("contact_person", "")) or "—"}</div>')
    h.append(f'<div class="kv"><b>Телефон:</b> '
             f'{_esc(r.get("contact_phone", "")) or "—"}</div>')
    h.append(f'<div class="kv"><b>Время работы:</b> '
             f'{_esc(r.get("work_hours", "")) or "—"}</div>')
    h.append(f'<div class="kv"><b>Перерыв:</b> {break_str}</div>')

    h.append("<h2>Особые условия</h2>")
    h.append(f'<div class="check"><span class="mark">{call_mark}</span>'
             f'Позвонить за 30 минут до приезда</div>')
    h.append(f'<div class="check"><span class="mark">{pass_mark}</span>'
             f'Нужен пропуск на территорию</div>')

    if r.get("comment"):
        h.append("<h2>Комментарий для водителя</h2>")
        h.append(f'<div class="comment">{_esc(r["comment"])}</div>')

    h.append("<h2>Состав груза</h2>")
    h.append("<table>")
    h.append("<tr>"
             "<th>Модель</th>"
             "<th class='num'>Кол-во</th>"
             "<th class='num'>Вес 1 шт, кг</th>"
             "<th class='num'>Общий вес, кг</th>"
             "</tr>")
    h.append(items_html)
    if total_weight_row:
        h.append(total_weight_row)
    h.append("</table>")

    if cargo_html:
        h.append("<h2>Грузовые характеристики</h2>")
        h.append(cargo_html)

    if r.get("tracking_number"):
        h.append(f'<p class="meta" style="margin-top:14px">'
                 f'Номер накладной ТК: <b>{_esc(r["tracking_number"])}</b></p>')

    h.append('<p class="meta" style="margin-top:20px">'
             'Сгенерировано программой PalletPacker</p>')

    return "".join(h)


# ============================================================
#                    PDF внутренний
# ============================================================

def export_internal_pdf(record: dict, path: str) -> None:
    html = _build_internal_html(record)
    _print_html(html, path)


def _build_internal_html(r: dict) -> str:
    entity_type = "Юр. лицо" if r.get("entity_type") == "legal" \
        else "Физ. лицо"

    break_str = "—"
    if r.get("has_break"):
        break_str = r.get("break_time", "") or "есть"

    def _y(v):
        return "да" if v else "нет"

    items_html = ""
    total_weight = 0.0
    total_qty = 0
    for it in r.get("items", []):
        model = _esc(it.get("model", ""))
        qty = int(it.get("qty", 0))
        total_qty += qty
        dims_txt = shp.format_dimensions(it.get("dims_mm")) or "—"
        try:
            w1 = float(it.get("weight_kg", 0) or 0)
        except (TypeError, ValueError):
            w1 = 0.0
        w_total = qty * w1
        total_weight += w_total
        serials = it.get("serials", []) or []
        serials_str = ", ".join(_esc(s) for s in serials) if serials else "—"

        items_html += (
            f"<tr>"
            f"<td>{model}</td>"
            f"<td class='num'>{qty}</td>"
            f"<td>{_esc(dims_txt)}</td>"
            f"<td class='num'>{w1:g}</td>"
            f"<td class='num'>{w_total:.2f}</td>"
            f"<td>{serials_str}</td>"
            f"</tr>"
        )
    if not items_html:
        items_html = ("<tr><td colspan='6' style='text-align:center;"
                      "color:#888'>нет позиций</td></tr>")

    # грузовые характеристики
    try:
        cargo_w = float(r.get("cargo_weight_kg", 0) or 0)
    except (TypeError, ValueError):
        cargo_w = 0
    try:
        cargo_v = float(r.get("cargo_volume_m3", 0) or 0)
    except (TypeError, ValueError):
        cargo_v = 0
    try:
        cargo_b = float(r.get("cargo_bbox_volume_m3", 0) or 0)
    except (TypeError, ValueError):
        cargo_b = 0
    try:
        cargo_p = int(r.get("cargo_pallets", 0) or 0)
    except (TypeError, ValueError):
        cargo_p = 0

    cargo_rows = []
    if total_weight:
        cargo_rows.append(("Вес груза (из состава), кг",
                           f"{total_weight:.2f}"))
    if cargo_w:
        cargo_rows.append(("Вес груза (в отчёте), кг", f"{cargo_w:.2f}"))
    if cargo_v:
        cargo_rows.append(("Объём коробок, м³", f"{cargo_v:.4f}"))
    if cargo_b:
        cargo_rows.append(("Габаритный объём, м³", f"{cargo_b:.4f}"))
    if cargo_p:
        cargo_rows.append(("Количество паллет, шт", str(cargo_p)))

    cargo_html = ""
    if cargo_rows:
        cargo_html = ['<table class="cargo-table">']
        for k, v in cargo_rows:
            cargo_html.append(
                f'<tr><td>{_esc(k)}</td><td class="num">{_esc(v)}</td></tr>'
            )
        cargo_html.append("</table>")
        cargo_html = "".join(cargo_html)

    status = shp.status_label(r)
    status_color = shp.status_color(r)

    css = """
    <style>
      body { font-family: Arial, sans-serif; color: #111; }
      h1 { color: #2e7d32; font-size: 17pt; margin: 0 0 4px 0; }
      h2 { font-size: 13pt; color: #333; margin: 14px 0 6px 0;
           border-bottom: 1px solid #ccc; padding-bottom: 2px; }
      .meta { color: #666; font-size: 10pt; }
      table { border-collapse: collapse; width: 100%; margin-top: 6px; }
      td, th { border: 1px solid #888; padding: 5px 8px; font-size: 9.5pt;
               vertical-align: top; }
      th { background: #e8f5e9; text-align: left; }
      .num { text-align: right; }
      .kv { margin: 3px 0; font-size: 10pt; }
      .kv b { display: inline-block; min-width: 200px; }
      .status-badge { display: inline-block; padding: 4px 12px;
                      border-radius: 12px; color: white; font-weight: bold;
                      font-size: 11pt; }
      .comment { background: #fffde7; border-left: 4px solid #fbc02d;
                 padding: 8px 12px; margin-top: 6px; font-size: 10pt; }
      .cargo-table { width: auto; min-width: 60%; }
      .cargo-table td { padding: 4px 10px; font-size: 10.5pt; }
      .cargo-table td:first-child { background: #f5f5f5; font-weight: bold; }
    </style>
    """

    h = [css]
    h.append('<h1>Учёт отправок ТК. Для внутренних нужд</h1>')
    h.append(
        f'<p class="meta">'
        f'Отправка № <b>{_esc(r.get("id", ""))}</b>'
        f' &nbsp;·&nbsp; '
        f'<span class="status-badge" style="background:{status_color}">'
        f'{status}</span>'
        f' &nbsp;·&nbsp; Создано: {_esc(r.get("created_at", ""))}'
        f'</p>'
    )

    h.append("<h2>Данные для транспортной компании</h2>")
    h.append(f'<div class="kv"><b>Адрес доставки:</b><br>'
             f'{_esc(r.get("address", "")) or "—"}</div>')
    h.append(f'<div class="kv"><b>Тип лица:</b> {entity_type}</div>')
    h.append(f'<div class="kv"><b>Наименование / ФИО:</b> '
             f'{_esc(r.get("entity_name", "")) or "—"}</div>')
    if r.get("entity_type") == "legal":
        h.append(f'<div class="kv"><b>ИНН:</b> '
                 f'{_esc(r.get("inn", "")) or "—"}</div>')
    h.append(f'<div class="kv"><b>Контактное лицо:</b> '
             f'{_esc(r.get("contact_person", "")) or "—"}</div>')
    h.append(f'<div class="kv"><b>Телефон:</b> '
             f'{_esc(r.get("contact_phone", "")) or "—"}</div>')
    h.append(f'<div class="kv"><b>Время работы:</b> '
             f'{_esc(r.get("work_hours", "")) or "—"}</div>')
    h.append(f'<div class="kv"><b>Перерыв:</b> {break_str}</div>')
    h.append(f'<div class="kv"><b>Позвонить за 30 мин:</b> '
             f'{_y(r.get("call_before_30min"))}</div>')
    h.append(f'<div class="kv"><b>Нужен пропуск:</b> '
             f'{_y(r.get("need_pass"))}</div>')
    if r.get("comment"):
        h.append(f'<div class="comment">{_esc(r["comment"])}</div>')

    h.append("<h2>Внутренний учёт</h2>")
    h.append(f'<div class="kv"><b>Дата отправки:</b> '
             f'{_esc(r.get("ship_date", "")) or "—"}</div>')
    h.append(f'<div class="kv"><b>Планируемый возврат:</b> '
             f'{_esc(r.get("planned_return_date", "")) or "—"}</div>')
    h.append(f'<div class="kv"><b>Фактический возврат:</b> '
             f'{_esc(r.get("actual_return_date", "")) or "—"}</div>')
    h.append(f'<div class="kv"><b>Кто отправил:</b> '
             f'{_esc(r.get("shipped_by", "")) or "—"}</div>')
    h.append(f'<div class="kv"><b>Номер накладной ТК:</b> '
             f'{_esc(r.get("tracking_number", "")) or "—"}</div>')
    try:
        cost = float(r.get("delivery_cost", 0) or 0)
    except (TypeError, ValueError):
        cost = 0
    h.append(f'<div class="kv"><b>Стоимость доставки:</b> {cost:.2f} ₽</div>')

    h.append("<h2>Состав отправки</h2>")
    h.append("<table>")
    h.append(
        "<tr>"
        "<th>Модель</th>"
        "<th class='num'>Кол-во</th>"
        "<th>Размер (Ш×Г×В), мм</th>"
        "<th class='num'>Вес 1 шт, кг</th>"
        "<th class='num'>Общий вес, кг</th>"
        "<th>Серийные номера</th>"
        "</tr>"
    )
    h.append(items_html)
    h.append("</table>")

    h.append(f'<div class="kv" style="margin-top:8px">'
             f'<b>Итого:</b> {total_qty} шт., '
             f'общий вес {total_weight:.2f} кг</div>')

    if cargo_html:
        h.append("<h2>Грузовые характеристики</h2>")
        h.append(cargo_html)

    h.append('<p class="meta" style="margin-top:20px">'
             'Сгенерировано программой PalletPacker</p>')

    return "".join(h)


# ============================================================
#                       Excel
# ============================================================

def export_shipments_xlsx(records: list, path: str) -> None:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()

    title_font = Font(bold=True, size=14, color="2E7D32")
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1976D2")
    total_font = Font(bold=True)
    total_fill = PatternFill("solid", fgColor="F1F8E9")

    # ---------- Лист 1: Отправки ----------
    ws = wb.active
    ws.title = "Отправки"
    ws["A1"] = "Отгрузки демо-оборудования"
    ws["A1"].font = title_font
    ws.merge_cells("A1:T1")
    ws["A2"] = (f"Всего: {len(records)} записей. "
                f"Экспорт от {datetime.now().strftime('%d.%m.%Y %H:%M')}")
    ws["A2"].font = Font(italic=True, size=10, color="666666")

    headers = [
        "Номер", "Статус", "Создано",
        "Адрес", "Тип лица", "Наименование/ФИО", "ИНН",
        "Контактное лицо", "Телефон", "Время работы", "Перерыв",
        "Позвонить за 30 мин", "Нужен пропуск", "Комментарий",
        "Дата отправки", "План. возврат", "Факт. возврат",
        "Кто отправил", "Номер накладной", "Стоимость доставки",
    ]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=c, value=h)
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    row = 5
    for r in records:
        break_str = "—"
        if r.get("has_break"):
            break_str = r.get("break_time", "") or "есть"
        vals = [
            r.get("id", ""),
            shp.status_label(r),
            r.get("created_at", ""),
            r.get("address", ""),
            "Юр. лицо" if r.get("entity_type") == "legal" else "Физ. лицо",
            r.get("entity_name", ""),
            r.get("inn", ""),
            r.get("contact_person", ""),
            r.get("contact_phone", ""),
            r.get("work_hours", ""),
            break_str,
            "да" if r.get("call_before_30min") else "нет",
            "да" if r.get("need_pass") else "нет",
            r.get("comment", ""),
            r.get("ship_date", ""),
            r.get("planned_return_date", ""),
            r.get("actual_return_date", ""),
            r.get("shipped_by", ""),
            r.get("tracking_number", ""),
            r.get("delivery_cost", 0),
        ]
        for c, v in enumerate(vals, 1):
            ws.cell(row=row, column=c, value=v)
        row += 1

    widths = [14, 14, 16, 30, 10, 26, 14, 22, 18, 14, 14,
              12, 12, 30, 13, 13, 13, 16, 18, 12]
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w

    # ---------- Лист 2: Состав (расширенный) ----------
    ws2 = wb.create_sheet("Состав")
    ws2["A1"] = "Состав отправок (детально)"
    ws2["A1"].font = title_font
    ws2.merge_cells("A1:I1")

    headers2 = [
        "Номер отправки", "Модель", "Кол-во",
        "Ш, мм", "Г, мм", "В, мм",
        "Вес 1 шт, кг", "Общий вес, кг", "Серийные номера",
    ]
    for c, h in enumerate(headers2, 1):
        cell = ws2.cell(row=3, column=c, value=h)
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    row = 4
    for r in records:
        for it in r.get("items", []):
            qty = int(it.get("qty", 0))
            dims = it.get("dims_mm") or [0, 0, 0]
            try:
                w1 = float(it.get("weight_kg", 0) or 0)
            except (TypeError, ValueError):
                w1 = 0.0
            w_total = qty * w1
            serials = it.get("serials", []) or []

            def _d(i):
                try:
                    v = float(dims[i]) if i < len(dims) else 0
                    return int(v) if v else ""
                except (TypeError, ValueError, IndexError):
                    return ""

            ws2.cell(row=row, column=1, value=r.get("id", ""))
            ws2.cell(row=row, column=2, value=it.get("model", ""))
            ws2.cell(row=row, column=3, value=qty)
            ws2.cell(row=row, column=4, value=_d(0))
            ws2.cell(row=row, column=5, value=_d(1))
            ws2.cell(row=row, column=6, value=_d(2))
            ws2.cell(row=row, column=7, value=w1 if w1 else "")
            ws2.cell(row=row, column=8, value=round(w_total, 3) if w_total else "")
            ws2.cell(row=row, column=9, value="; ".join(serials))
            row += 1

    widths2 = [16, 32, 10, 9, 9, 9, 12, 14, 40]
    for c, w in enumerate(widths2, 1):
        ws2.column_dimensions[get_column_letter(c)].width = w

    # ---------- Лист 3: Сводка ----------
    ws3 = wb.create_sheet("Сводка")
    ws3["A1"] = "Сводка по отправкам"
    ws3["A1"].font = title_font
    ws3.merge_cells("A1:J1")

    headers3 = [
        "Номер", "Создано", "Статус", "Получатель",
        "Позиций", "Коробок", "Общий вес, кг",
        "Объём, м³", "Паллет", "Дата отправки",
    ]
    for c, h in enumerate(headers3, 1):
        cell = ws3.cell(row=3, column=c, value=h)
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    row = 4
    for r in records:
        items = r.get("items", [])
        total_qty = sum(int(it.get("qty", 0)) for it in items)
        weight = shp.calc_weight_from_items(items)
        volume = shp.calc_volume_from_items(items)
        try:
            pallets = int(r.get("cargo_pallets", 0) or 0)
        except (TypeError, ValueError):
            pallets = 0

        ws3.cell(row=row, column=1, value=r.get("id", ""))
        ws3.cell(row=row, column=2, value=r.get("created_at", ""))
        ws3.cell(row=row, column=3, value=shp.status_label(r))
        ws3.cell(row=row, column=4, value=r.get("entity_name", ""))
        ws3.cell(row=row, column=5, value=len(items))
        ws3.cell(row=row, column=6, value=total_qty)
        ws3.cell(row=row, column=7,
                 value=round(weight, 2) if weight else 0)
        ws3.cell(row=row, column=8,
                 value=round(volume, 4) if volume else 0)
        ws3.cell(row=row, column=9, value=pallets)
        ws3.cell(row=row, column=10, value=r.get("ship_date", ""))
        row += 1

    widths3 = [16, 16, 14, 26, 10, 10, 14, 12, 10, 14]
    for c, w in enumerate(widths3, 1):
        ws3.column_dimensions[get_column_letter(c)].width = w

    wb.save(path)


def import_shipments_xlsx(path: str) -> tuple:
    import openpyxl

    warnings = []
    records = []

    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:
        raise RuntimeError(f"Не удалось открыть файл: {e}")

    if "Отправки" not in wb.sheetnames:
        raise RuntimeError(
            "В файле нет листа «Отправки». "
            "Это должен быть файл, созданный через «📤 Excel экспорт»."
        )

    ws = wb["Отправки"]

    header_row = None
    for r in range(1, 15):
        v = ws.cell(row=r, column=1).value
        if v and str(v).strip().lower() in ("номер", "id"):
            header_row = r
            break
    if header_row is None:
        raise RuntimeError("Не найдена строка заголовков.")

    cols = {}
    for c in range(1, 40):
        v = ws.cell(row=header_row, column=c).value
        if v:
            cols[str(v).strip().lower()] = c

    # --- состав с расширенными колонками ---
    items_map = {}
    if "Состав" in wb.sheetnames:
        ws2 = wb["Состав"]
        # ищем строку заголовков
        hdr2 = None
        for r in range(1, 8):
            v = ws2.cell(row=r, column=1).value
            if v and str(v).strip().lower() in ("номер отправки",):
                hdr2 = r
                break
        if hdr2 is None:
            hdr2 = 3

        cols2 = {}
        for c in range(1, 20):
            v = ws2.cell(row=hdr2, column=c).value
            if v:
                cols2[str(v).strip().lower()] = c

        def _g2(row_num, key):
            c = cols2.get(key)
            if c is None:
                return ""
            v = ws2.cell(row=row_num, column=c).value
            return "" if v is None else str(v).strip()

        for r in range(hdr2 + 1, ws2.max_row + 1):
            sid = ws2.cell(row=r, column=1).value
            if not sid:
                continue
            sid = str(sid).strip()

            model = _g2(r, "модель")
            if not model:
                continue
            try:
                qty = int(float(_g2(r, "кол-во") or 0))
            except ValueError:
                qty = 0

            def _f2(key):
                s = _g2(r, key)
                try:
                    return float(s.replace(",", ".") or 0)
                except ValueError:
                    return 0.0

            dims = [
                _f2("ш, мм"),
                _f2("г, мм"),
                _f2("в, мм"),
            ]
            weight = _f2("вес 1 шт, кг")
            serials_raw = _g2(r, "серийные номера")
            serials = [s.strip() for s in
                       serials_raw.replace(",", ";").split(";")
                       if s.strip()]

            items_map.setdefault(sid, []).append({
                "model": model,
                "qty": qty,
                "dims_mm": dims,
                "weight_kg": weight,
                "serials": serials,
            })

    def _get(row_num, key):
        c = cols.get(key)
        if c is None:
            return ""
        v = ws.cell(row=row_num, column=c).value
        return "" if v is None else str(v).strip()

    def _get_float(row_num, key):
        s = _get(row_num, key)
        try:
            return float(s.replace(",", ".") or 0)
        except ValueError:
            return 0.0

    def _get_int(row_num, key):
        s = _get(row_num, key)
        try:
            return int(float(s.replace(",", ".") or 0))
        except ValueError:
            return 0

    for r in range(header_row + 1, ws.max_row + 1):
        sid = _get(r, "номер")
        if not sid:
            continue

        entity_str = _get(r, "тип лица").lower()
        entity_type = "physical" if "физ" in entity_str else "legal"

        has_break_str = _get(r, "перерыв").lower()
        has_break = has_break_str not in ("", "—", "нет", "-")
        break_time = ""
        if has_break and has_break_str not in ("есть", "да"):
            break_time = _get(r, "перерыв")

        call_str = _get(r, "позвонить за 30 мин").lower()
        need_pass_str = _get(r, "нужен пропуск").lower()

        rec = {
            "id": sid,
            "created_at": _get(r, "создано"),
            "status": "in_transit",
            "address": _get(r, "адрес"),
            "entity_type": entity_type,
            "entity_name": _get(r, "наименование/фио"),
            "inn": _get(r, "инн"),
            "contact_person": _get(r, "контактное лицо"),
            "contact_phone": _get(r, "телефон"),
            "work_hours": _get(r, "время работы"),
            "has_break": has_break,
            "break_time": break_time,
            "call_before_30min": call_str.startswith("да"),
            "need_pass": need_pass_str.startswith("да"),
            "comment": _get(r, "комментарий"),
            "ship_date": _get(r, "дата отправки"),
            "planned_return_date": _get(r, "план. возврат"),
            "actual_return_date": _get(r, "факт. возврат"),
            "shipped_by": _get(r, "кто отправил"),
            "tracking_number": _get(r, "номер накладной"),
            "delivery_cost": _get_float(r, "стоимость доставки"),
            "items": items_map.get(sid, []),
        }

        # грузовые характеристики: если в файле нет колонок — считаем из состава
        rec["cargo_weight_kg"] = _get_float(r, "вес груза, кг")
        rec["cargo_volume_m3"] = _get_float(r, "объём коробок, м³")
        rec["cargo_bbox_volume_m3"] = _get_float(r, "габаритный объём, м³")
        rec["cargo_pallets"] = _get_int(r, "кол-во паллет")

        # если в файле этих колонок не было — оценим из состава
        if not rec["cargo_weight_kg"]:
            rec["cargo_weight_kg"] = round(
                shp.calc_weight_from_items(rec["items"]), 2
            )
        if not rec["cargo_volume_m3"]:
            rec["cargo_volume_m3"] = round(
                shp.calc_volume_from_items(rec["items"]), 4
            )

        records.append(rec)

    return records, warnings


# ============================================================
#                       helpers
# ============================================================

def _esc(s) -> str:
    if s is None:
        return ""
    s = str(s)
    return (s.replace("&", "&amp;").replace("<", "&lt;")
             .replace(">", "&gt;"))


def _print_html(html: str, path: str) -> None:
    doc = QTextDocument()
    doc.setHtml(html)

    printer = QPrinter(QPrinter.HighResolution)
    printer.setOutputFormat(QPrinter.PdfFormat)
    printer.setOutputFileName(path)
    printer.setPageSize(QPageSize(QPageSize.A4))
    printer.setPageMargins(QMarginsF(15, 15, 15, 15))
    doc.print_(printer)