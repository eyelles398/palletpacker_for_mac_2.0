"""
Экспорт текущего заказа в Excel.

Формат полностью совместим с order_import.py:
    A — модель
    B — кол-во
    C — паллета №

Сверху добавляется шапка с параметрами паллеты, заказчиком, датой.
Импорт эти строки молча пропускает.
"""
from __future__ import annotations

from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment


def export_order_xlsx(items, path, pallet_params=None,
                      customer="", spec_number=""):
    """
    items: [{"name": str, "qty": int, "pallet_group": int}, ...]
    pallet_params: {"L": float, "W": float, "H": float, "max_weight": float}
                   (опционально)
    customer, spec_number: необязательные строки для шапки.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Заказ"

    title_font = Font(bold=True, size=14, color="2E7D32")
    head_font = Font(bold=True)
    total_font = Font(bold=True, color="1B5E20")
    header_fill = PatternFill("solid", fgColor="E8F5E9")
    total_fill = PatternFill("solid", fgColor="F1F8E9")
    small_font = Font(size=10, color="666666")

    # ---- шапка ----
    ws["A1"] = "Заказ на укладку коробок"
    ws["A1"].font = title_font
    ws.merge_cells("A1:C1")

    row = 2
    if pallet_params:
        L = pallet_params.get("L", 0)
        W = pallet_params.get("W", 0)
        H = pallet_params.get("H", 0)
        MW = pallet_params.get("max_weight", 0)
        ws.cell(
            row=row, column=1,
            value=f"Паллета: {L:.0f}×{W:.0f}×{H:.0f} мм, до {MW:.0f} кг"
        ).font = small_font
        row += 1

    if customer:
        ws.cell(row=row, column=1,
                value=f"Заказчик: {customer}").font = small_font
        row += 1
    if spec_number:
        ws.cell(row=row, column=1,
                value=f"Спецификация: {spec_number}").font = small_font
        row += 1

    ws.cell(
        row=row, column=1,
        value=f"Дата: {datetime.now().strftime('%d.%m.%Y %H:%M')}"
    ).font = small_font
    row += 1

    row += 1  # пустая строка перед таблицей

    # ---- заголовки таблицы ----
    headers = ["Модель", "Кол-во", "Паллета №"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = head_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    row += 1

    # ---- данные ----
    total_qty = 0
    for it in items:
        name = it.get("name", "")
        qty = int(it.get("qty", 0))
        pg = int(it.get("pallet_group", 1) or 1)

        ws.cell(row=row, column=1, value=name)
        c_qty = ws.cell(row=row, column=2, value=qty)
        c_qty.alignment = Alignment(horizontal="right")
        c_pg = ws.cell(row=row, column=3, value=pg)
        c_pg.alignment = Alignment(horizontal="right")

        total_qty += qty
        row += 1

    # ---- итоговая строка ----
    if items:
        row += 1
        ws.cell(
            row=row, column=1,
            value=f"ИТОГО: {len(items)} позиций, {total_qty} коробок"
        )
        ws.cell(row=row, column=1).font = total_font
        for c in (1, 2, 3):
            ws.cell(row=row, column=c).fill = total_fill

    # ---- ширины колонок ----
    ws.column_dimensions["A"].width = 45
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 14

    wb.save(path)